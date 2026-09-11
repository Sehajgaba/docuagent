"""Build docs/docuagent-handbook.docx from handbook_content.py.

Usage:
    python scripts/build_handbook.py
    python scripts/build_handbook.py --out docs/other-name.docx

The .docx is a generated artifact and is gitignored, matching this repo's
existing convention (see .gitignore: generated outputs stay local, sources are
committed). Re-run this script to regenerate it.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document  # noqa: E402
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

import handbook_content as content  # noqa: E402
from handbook_blocks import Bullets, Callout, Code, H, P, QA, Table  # noqa: E402

CODE_BG = "F2F2F2"
CALLOUT_BG = {"war": "FFF4E5", "note": "EAF2FB", "gotcha": "FDECEC"}
CALLOUT_LABEL = {"war": "THE REAL FINDING", "note": "NOTE", "gotcha": "GOTCHA"}


# --- low-level helpers -------------------------------------------------------
def _shade(paragraph, fill: str) -> None:
    """Apply a background fill to a paragraph (python-docx has no API for this)."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def _field(paragraph, instruction: str) -> None:
    """Insert a Word field code (used for the TOC and page numbers)."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, end):
        run._r.append(element)


def _define_styles(doc: Document) -> None:
    code = doc.styles.add_style("HandbookCode", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    code.paragraph_format.space_before = Pt(6)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.left_indent = Inches(0.25)

    callout = doc.styles.add_style("HandbookCallout", WD_STYLE_TYPE.PARAGRAPH)
    callout.font.size = Pt(10)
    callout.paragraph_format.space_before = Pt(8)
    callout.paragraph_format.space_after = Pt(8)
    callout.paragraph_format.left_indent = Inches(0.2)
    callout.paragraph_format.right_indent = Inches(0.2)

    answer = doc.styles.add_style("HandbookAnswer", WD_STYLE_TYPE.PARAGRAPH)
    answer.font.size = Pt(10.5)
    answer.paragraph_format.left_indent = Inches(0.3)
    answer.paragraph_format.space_after = Pt(10)


def _footer_page_numbers(doc: Document) -> None:
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(footer, "PAGE")


# --- block renderers ---------------------------------------------------------
def _render_code(doc: Document, block: Code) -> None:
    for line in block.text.split("\n"):
        paragraph = doc.add_paragraph(line or " ", style="HandbookCode")
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        _shade(paragraph, CODE_BG)


def _render_table(doc: Document, block: Table) -> None:
    table = doc.add_table(rows=1, cols=len(block.headers))
    table.style = "Table Grid"  # always present in the default template
    for cell, heading in zip(table.rows[0].cells, block.headers):
        cell.text = ""
        run = cell.paragraphs[0].add_run(heading)
        run.bold = True
        run.font.size = Pt(9.5)
    for row_values in block.rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            run.font.size = Pt(9.5)
    doc.add_paragraph()


def _render_callout(doc: Document, block: Callout) -> None:
    paragraph = doc.add_paragraph(style="HandbookCallout")
    label = paragraph.add_run(f"{CALLOUT_LABEL.get(block.kind, 'NOTE')}  ")
    label.bold = True
    label.font.size = Pt(9)
    text = block.text
    # Content already leads with the label for war stories; avoid printing twice.
    for prefix in ("THE REAL FINDING — ", "THE OPEN PROBLEM THIS ADDRESSES — "):
        if text.startswith(prefix):
            label.text = prefix.replace(" — ", "  ")
            text = text[len(prefix):]
            break
    paragraph.add_run(text)
    _shade(paragraph, CALLOUT_BG.get(block.kind, "EAF2FB"))


def _render_qa(doc: Document, block: QA) -> None:
    question = doc.add_paragraph()
    question.paragraph_format.space_after = Pt(2)
    run = question.add_run(f"Q: {block.question}")
    run.bold = True
    run.font.size = Pt(10.5)
    answer = doc.add_paragraph(block.answer, style="HandbookAnswer")
    answer.paragraph_format.left_indent = Inches(0.3)


def _render_block(doc: Document, block, heading_shift: int = 0) -> None:
    if isinstance(block, H):
        doc.add_heading(block.text, level=min(block.level + heading_shift, 9))
    elif isinstance(block, P):
        doc.add_paragraph(block.text)
    elif isinstance(block, Bullets):
        for item in block.items:
            doc.add_paragraph(item, style="List Bullet")
    elif isinstance(block, Code):
        _render_code(doc, block)
    elif isinstance(block, Table):
        _render_table(doc, block)
    elif isinstance(block, QA):
        _render_qa(doc, block)
    elif isinstance(block, Callout):
        _render_callout(doc, block)
    else:
        raise TypeError(f"unknown block type: {type(block).__name__}")


# --- document assembly -------------------------------------------------------
def _front_matter(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(content.TITLE)
    run.bold = True
    run.font.size = Pt(26)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(content.SUBTITLE)
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    byline = doc.add_paragraph()
    byline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    byline.add_run(f"{content.OWNER} · generated from the build log").font.size = Pt(9)

    doc.add_paragraph()
    doc.add_heading("How to use this", level=2)
    for item in content.HOW_TO_USE:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_page_break()
    doc.add_heading("Contents", level=1)
    toc = doc.add_paragraph()
    _field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    note = doc.add_paragraph()
    run = note.add_run(
        "If the contents list above is empty, click it and press F9 (or right-click "
        "and choose Update Field) — Word builds it on demand."
    )
    run.italic = True
    run.font.size = Pt(9)
    doc.add_page_break()


def _render_concept(doc: Document, concept) -> None:
    doc.add_heading(f"{concept.number}. {concept.title}", level=3)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(10)
    status = meta.add_run(concept.status)
    status.bold = True
    status.font.size = Pt(8.5)
    status.font.color.rgb = (
        RGBColor(0x1F, 0x6F, 0x3C) if concept.status == "BUILT" else RGBColor(0x99, 0x55, 0x00)
    )
    tail = meta.add_run(f"   ·   {concept.day}")
    tail.font.size = Pt(8.5)
    tail.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # Concept headings sit at level 3, so sub-headings inside one shift to 4.
    for block in concept.blocks:
        _render_block(doc, block, heading_shift=1)


def build(out_path: Path) -> Document:
    doc = Document()
    _define_styles(doc)
    _footer_page_numbers(doc)
    _front_matter(doc)

    for part in content.PARTS:
        doc.add_heading(part.title, level=1)
        if part.intro:
            intro = doc.add_paragraph()
            run = intro.add_run(part.intro)
            run.italic = True
        for section in part.sections:
            doc.add_heading(section.title, level=2)
            if section.intro:
                intro = doc.add_paragraph()
                run = intro.add_run(section.intro)
                run.italic = True
            for block in section.blocks:
                _render_block(doc, block)
            for concept in section.concepts:
                _render_concept(doc, concept)
        doc.add_page_break()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return doc


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--out",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "docs" / "docuagent-handbook.docx",
    )
    args = ap.parse_args()

    build(args.out)
    size_kb = args.out.stat().st_size / 1024
    concepts = sum(len(s.concepts) for p in content.PARTS for s in p.sections)
    print(f"-> wrote {args.out}  ({size_kb:,.0f} KB, {concepts} concepts)")


if __name__ == "__main__":
    main()
